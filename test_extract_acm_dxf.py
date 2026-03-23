#!/usr/bin/env python3
import argparse
import json
import re
from pathlib import Path
from typing import List, Dict, Any, Optional, Union

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment


PREFERRED_LAYERS = ['CUT', 'PROFILE', 'OUTLINE', 'PART', 'PANEL', 'ROUT']
BAD_LAYERS = {'SHEET', 'BORDER', 'TITLE'}


def safe_float(v: str):
    try:
        return float(v)
    except Exception:
        return None


def parse_ascii_dxf_entities(path: Path) -> List[Dict[str, Any]]:
    text = path.read_text(encoding='latin1')
    lines = text.splitlines()
    pairs = []
    i = 0
    while i < len(lines) - 1:
        code = lines[i].strip()
        value = lines[i + 1].rstrip('\n')
        pairs.append((code, value))
        i += 2

    in_entities = False
    entities = []
    current = None
    j = 0
    while j < len(pairs):
        code, value = pairs[j]
        if code == '0' and value == 'SECTION':
            if j + 1 < len(pairs) and pairs[j + 1][0] == '2' and pairs[j + 1][1] == 'ENTITIES':
                in_entities = True
                j += 2
                continue
        if in_entities and code == '0' and value == 'ENDSEC':
            break
        if not in_entities:
            j += 1
            continue
        if code == '0':
            if current:
                entities.append(current)
            current = {'type': value, 'tags': []}
        else:
            if current is not None:
                current['tags'].append((code, value))
        j += 1
    if current:
        entities.append(current)
    return entities


def entity_layer(ent: Dict[str, Any]) -> str:
    for c, v in ent['tags']:
        if c == '8':
            return v.strip()
    return ''


def plain_mtext(raw: str) -> str:
    s = raw.replace('\\P', '\n')
    s = re.sub(r'\\[A-Za-z][^;]*;?', '', s)
    s = s.replace('{', '').replace('}', '')
    return s.strip()


def extract_text_entities(entities: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    items = []
    for ent in entities:
        et = ent['type']
        tags = ent['tags']
        if et == 'MTEXT':
            content = ''.join(v for c, v in tags if c in ('1', '3'))
            x = next((safe_float(v) for c, v in tags if c == '10'), None)
            y = next((safe_float(v) for c, v in tags if c == '20'), None)
            items.append({
                'entity': et,
                'layer': entity_layer(ent),
                'text': plain_mtext(content),
                'insert': [x, y, 0.0],
            })
        elif et == 'TEXT':
            content = next((v for c, v in tags if c == '1'), '').strip()
            x = next((safe_float(v) for c, v in tags if c == '10'), None)
            y = next((safe_float(v) for c, v in tags if c == '20'), None)
            items.append({
                'entity': et,
                'layer': entity_layer(ent),
                'text': content,
                'insert': [x, y, 0.0],
            })
    return items


def extract_lwpolylines(entities: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    polys = []
    for ent in entities:
        if ent['type'] != 'LWPOLYLINE':
            continue
        tags = ent['tags']
        layer = entity_layer(ent)
        closed = any(c == '70' and (int(float(v)) & 1) == 1 for c, v in tags if v.strip())

        pts = []
        i = 0
        while i < len(tags):
            c, v = tags[i]
            if c == '10':
                x = safe_float(v)
                y = None
                k = i + 1
                while k < len(tags):
                    c2, v2 = tags[k]
                    if c2 == '20':
                        y = safe_float(v2)
                        break
                    if c2 == '10':
                        break
                    k += 1
                if x is not None and y is not None:
                    pts.append((x, y))
            i += 1

        if not pts:
            continue

        xs = [p[0] for p in pts]
        ys = [p[1] for p in pts]

        perim = 0.0
        for idx in range(len(pts) - 1):
            x1, y1 = pts[idx]
            x2, y2 = pts[idx + 1]
            perim += ((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5
        if closed and len(pts) > 1:
            x1, y1 = pts[-1]
            x2, y2 = pts[0]
            perim += ((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5

        polys.append({
            'entity': 'LWPOLYLINE',
            'layer': layer,
            'closed': closed,
            'point_count': len(pts),
            'points': pts,
            'bbox': {
                'min_x': min(xs),
                'min_y': min(ys),
                'max_x': max(xs),
                'max_y': max(ys),
                'width': max(xs) - min(xs),
                'height': max(ys) - min(ys),
            },
            'perimeter': perim,
        })
    return polys


def choose_best_text_meta(text_items: List[Dict[str, Any]]) -> Dict[str, Any]:
    candidates = [t for t in text_items if 'PANEL MARK' in t['text'].upper()]
    if not candidates:
        candidates = text_items[:]

    best = candidates[0] if candidates else None
    meta: Dict[str, Any] = {}

    if best:
        raw = best['text'].replace('\n', ' ')
        m = re.search(r'PANEL MARKS?\s*:\s*([A-Za-z0-9._-]+)', raw, re.I)
        if m:
            meta['tag'] = m.group(1).strip()

        m = re.search(r'QUANTITY\s*:\s*([0-9]+)', raw, re.I)
        if m:
            meta['qty'] = int(m.group(1))

        m = re.search(r'FINISH\s*:\s*(.+)$', raw, re.I)
        if m:
            meta['finish'] = m.group(1).strip()

        meta['source_text'] = best['text']

    if 'tag' not in meta:
        pat = re.compile(r'^[A-Za-z0-9._-]+$')
        for t in text_items:
            txt = t['text'].strip()
            if pat.match(txt):
                meta['tag'] = txt
                break

    if 'qty' not in meta:
        meta['qty'] = 1

    if 'finish' not in meta:
        meta['finish'] = None

    return meta


def choose_best_polyline(polys: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    closed = [p for p in polys if p['closed']]
    if not closed:
        return None

    def score(p: Dict[str, Any]) -> float:
        layer = p['layer'].upper()
        score = p['bbox']['width'] * p['bbox']['height']
        if layer in PREFERRED_LAYERS:
            score += 1_000_000
        if layer in BAD_LAYERS:
            score -= 1_000_000
        return score

    closed.sort(key=score, reverse=True)
    return closed[0]


def record_from_dxf(path: Path) -> Dict[str, Any]:
    entities = parse_ascii_dxf_entities(path)
    texts = extract_text_entities(entities)
    polys = extract_lwpolylines(entities)

    meta = choose_best_text_meta(texts)
    poly = choose_best_polyline(polys)

    rec = {
        'file': path.name,
        'tag': meta.get('tag'),
        'qty': meta.get('qty', 1),
        'stretch_in': round(poly['bbox']['width'], 2) if poly else None,
        'height_in': round(poly['bbox']['height'], 2) if poly else None,
        'sq_ft_each': round((poly['bbox']['width'] * poly['bbox']['height']) / 144.0, 2) if poly else None,
        'sq_ft_total': round(((poly['bbox']['width'] * poly['bbox']['height']) / 144.0) * meta.get('qty', 1), 2) if poly else None,
        'finish': meta.get('finish'),
        'geometry_layer': poly['layer'] if poly else None,
        'point_count': poly['point_count'] if poly else None,
        'perimeter_in': round(poly['perimeter'], 2) if poly else None,
        'source_text': meta.get('source_text'),
        'bbox': poly['bbox'] if poly else None,
        'review_flags': [],
    }

    if not rec['tag']:
        rec['review_flags'].append('missing_tag')
    if poly is None:
        rec['review_flags'].append('missing_closed_polyline')
    return rec


def build_workbook(records: List[Dict[str, Any]], out_xlsx: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Panel Schedule'

    ws.merge_cells('F1:G1')
    ws['F1'] = 'Sq Ft'

    headers = ['Tag', 'Qty', 'Stretch (in)', 'Height (in)', 'Finish', 'Each', 'Total']
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=2, column=idx).value = header

    fills = {
        'header': PatternFill('solid', fgColor='D9D9D9'),
        'subheader': PatternFill('solid', fgColor='E7E7E7'),
    }
    border = Border(
        left=Side(style='thin', color='B7B7B7'),
        right=Side(style='thin', color='B7B7B7'),
        top=Side(style='thin', color='B7B7B7'),
        bottom=Side(style='thin', color='B7B7B7'),
    )

    for cell in ws[1]:
        cell.fill = fills['header']
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    for cell in ws[2]:
        cell.fill = fills['subheader']
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    for col in range(1, 6):
        c = ws.cell(row=1, column=col)
        c.fill = fills['header']
        c.border = border

    widths = {
        'A': 16,
        'B': 10,
        'C': 16,
        'D': 16,
        'E': 28,
        'F': 12,
        'G': 12,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 28

    start_row = 3
    for r_idx, rec in enumerate(records, start=start_row):
        ws.cell(r_idx, 1).value = rec.get('tag')
        ws.cell(r_idx, 2).value = rec.get('qty')
        ws.cell(r_idx, 3).value = rec.get('stretch_in')
        ws.cell(r_idx, 4).value = rec.get('height_in')
        ws.cell(r_idx, 5).value = rec.get('finish')
        ws.cell(r_idx, 6).value = f'=ROUND((C{r_idx}*D{r_idx})/144,2)'
        ws.cell(r_idx, 7).value = f'=ROUND(B{r_idx}*F{r_idx},2)'

        for c in range(1, 8):
            cell = ws.cell(r_idx, c)
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            if c in (3, 4, 6, 7):
                cell.number_format = '0.00'
            elif c == 2:
                cell.number_format = '0'

    raw = wb.create_sheet('Raw Data')
    raw_headers = [
        'file', 'tag', 'qty', 'stretch_in', 'height_in', 'sq_ft_each', 'sq_ft_total',
        'finish', 'geometry_layer', 'point_count', 'perimeter_in', 'source_text', 'review_flags'
    ]
    for col_idx, h in enumerate(raw_headers, start=1):
        raw.cell(1, col_idx).value = h
        raw.cell(1, col_idx).font = Font(bold=True)

    for row_idx, rec in enumerate(records, start=2):
        for col_idx, h in enumerate(raw_headers, start=1):
            value = rec.get(h)
            if isinstance(value, list):
                value = ', '.join(map(str, value))
            raw.cell(row_idx, col_idx).value = value

    wb.save(out_xlsx)


def extract_panel_schedule(
    input_folder: Union[str, Path],
    json_out: str = 'panel_schedule.json',
    xlsx_out: str = 'panel_schedule.xlsx',
) -> Dict[str, Any]:
    """
    Importable function for FastAPI/backend use.

    Keeps the same logic as main():
    - reads all .dxf files in input_folder
    - writes JSON + XLSX into that same folder
    - returns useful metadata for the caller
    """
    input_folder = Path(input_folder)
    files = sorted(input_folder.glob('*.dxf'))
    if not files:
        raise ValueError(f'No .dxf files found in {input_folder}')

    records = [record_from_dxf(path) for path in files]

    json_out_path = input_folder / json_out
    xlsx_out_path = input_folder / xlsx_out

    json_out_path.write_text(json.dumps(records, indent=2), encoding='utf-8')
    build_workbook(records, xlsx_out_path)

    return {
        'records': records,
        'record_count': len(records),
        'json_out': str(json_out_path),
        'xlsx_out': str(xlsx_out_path),
    }


def main():
    parser = argparse.ArgumentParser(description='Extract panel schedule from ASCII DXF files.')
    parser.add_argument('input_folder', help='Folder containing .dxf files')
    parser.add_argument('--json-out', default='panel_schedule.json')
    parser.add_argument('--xlsx-out', default='panel_schedule.xlsx')
    args = parser.parse_args()

    result = extract_panel_schedule(
        input_folder=args.input_folder,
        json_out=args.json_out,
        xlsx_out=args.xlsx_out,
    )

    print(f"Wrote {result['json_out']}")
    print(f"Wrote {result['xlsx_out']}")


if __name__ == '__main__':
    main()