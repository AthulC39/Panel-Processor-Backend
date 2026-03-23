from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from test_extract_acm_dxf import extract_panel_schedule


def run_processing_pipeline(input_files, output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    dxf_files = [Path(p) for p in input_files if Path(p).suffix.lower() == ".dxf"]

    if not dxf_files:
        raise ValueError("No DXF files were uploaded or generated.")

    input_folder = dxf_files[0].parent

    temp_xlsx_path = output_dir / "panel_bom.xlsx"
    json_path = output_dir / "panel_schedule.json"

    result = extract_panel_schedule(
        input_folder=input_folder,
        json_out=str(json_path),
        xlsx_out=str(temp_xlsx_path),
    )

    workbook_path = Path(result["xlsx_out"])

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not created: {workbook_path}")

    wb = load_workbook(workbook_path)

    if not wb.sheetnames:
        raise ValueError("Generated workbook has no sheets.")

    raw_sheet_names = [name for name in wb.sheetnames if name.strip().lower() == "raw data"]
    for sheet_name in raw_sheet_names:
        if len(wb.sheetnames) > 1:
            ws = wb[sheet_name]
            wb.remove(ws)

    remaining_sheets = wb.sheetnames
    if not remaining_sheets:
        raise ValueError("Workbook has no sheets left after removing raw data sheet.")

    wb[remaining_sheets[0]].title = "BOM"

    date_str = datetime.now().strftime("%Y-%m-%d")
    final_workbook_path = output_dir / f"Panel_BOM_{date_str}.xlsx"

    wb.save(final_workbook_path)

    if workbook_path != final_workbook_path and workbook_path.exists():
        workbook_path.unlink()

    return {
        "output_file": str(final_workbook_path),
        "record_count": result.get("record_count", 0),
    }